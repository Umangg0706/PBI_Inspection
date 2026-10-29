#!/usr/bin/env python3
"""
PBIX Inspector v2 - drag a .pbix (or .pbit) onto a local web page and get an Excel
workbook with:

  Tables         every table + details
  Columns        Table, Folder, Column, Data Type, Is Calculated, Formula, Is Hidden
  Measures       Measure, Formula, Table, Folder
  Relationships  From/To table + column, cardinality, filter direction, active
  M Queries      every Power Query: source system, source object, steps, full M code
  Visual Design  Report, Page, Visual Type, Visual Title, every field used + where it is used
  Lineage        Visual -> Field -> Measure/Calc-column chain -> Base Table[Column] -> Source system/object

Auto-generated items (LocalDateTable_*, DateTableTemplate_*, RowNumber-* columns and
relationships to them) are left out. Everything runs locally; nothing is uploaded.

SETUP (once):    pip install -r requirements.txt
RUN (drag/drop): python pbix_inspector.py
RUN (CLI):       python pbix_inspector.py MyReport.pbix  ->  MyReport_inspection.xlsx
"""
import io
import json
import os
import re
import shutil
import sqlite3
import sys
import tempfile
import threading
import uuid
import webbrowser
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, unquote, urlparse

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------------
# Sheets (tab order)
# ----------------------------------------------------------------------------
SHEETS = {
    "Tables": ["Table Name", "Table Type", "Is Hidden", "Column Count",
               "Calculated Column Count", "Measure Count", "Storage Mode",
               "Has M Query", "Calculated Table DAX"],
    "Columns": ["Table Name", "Folder Name", "Column Name", "Data Type",
                "Is Calculated Column", "Formula", "Is Hidden"],
    "Measures": ["Measure Name", "Formula", "Table Name", "Folder Name"],
    "Relationships": ["From Table", "From Column", "To Table", "To Column",
                      "Cardinality", "Cross Filter Direction", "Is Active"],
    "M Queries": ["Table / Query Name", "Source System", "Server / Path",
                  "Source Object", "Transformation Steps", "M Expression"],
    "Visual Design": ["Report Name", "Page Name", "Visual Type", "Visual Title",
                      "Field Name", "Field Type", "Table Name", "Used In",
                      "Aggregation", "X", "Y", "Width", "Height", "Visual ID"],
    "Lineage": ["Report Name", "Page Name", "Visual Type", "Visual Title",
                "Field Used", "Field Type", "Field Table", "Field Formula",
                "Dependency Path (DAX)", "Model Table", "Model Column",
                "Power Query Transformation", "Source Table", "Source Column",
                "Source System", "Server / Path", "Source Object",
                "Full Lineage Path", "Source SQL / Query"],
    "Impact Analysis": ["Source System", "Source Object", "Source Table", "Source Column",
                        "Model Table", "Model Column", "Dependency Chain",
                        "Power Query Transformation", "Pages Affected",
                        "Visuals Affected (count)", "Affected Visuals"],
}

AUTO_TABLE = re.compile(r"^(LocalDateTable|DateTableTemplate)_", re.I)
AUTO_COLUMN = re.compile(r"^RowNumber-", re.I)
AGG = {0: "Sum", 1: "Average", 2: "Min", 3: "Max", 4: "Count", 5: "Count (Distinct)",
       6: "Median", 7: "Standard Deviation", 8: "Variance"}


def is_auto_table(n):
    return bool(AUTO_TABLE.match(n or ""))


def is_auto_column(n):
    return bool(AUTO_COLUMN.match(n or ""))


def clean(v):
    if v is None:
        return ""
    if hasattr(v, "item"):
        try:
            v = v.item()
        except Exception:
            pass
    if isinstance(v, float) and v != v:
        return ""
    if isinstance(v, (list, tuple)):
        return "\n".join(str(x) for x in v)
    return v


def join_expr(v):
    return "\n".join(str(x) for x in v) if isinstance(v, list) else (v or "")


def decode_bytes(raw):
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return raw.decode("utf-16").lstrip("\ufeff")
    if b"\x00" in raw[:20]:
        return raw.decode("utf-16-le").lstrip("\ufeff")
    return raw.decode("utf-8-sig")


def new_result():
    out = {name: [] for name in SHEETS}
    out["_visuals"] = []
    return out, []


# ----------------------------------------------------------------------------
# Model extraction A: DataModelSchema JSON (.pbit)
# ----------------------------------------------------------------------------
def from_schema(schema, out):
    model = schema.get("model", schema)
    for t in model.get("tables", []):
        tname = t.get("name", "")
        annotations = {a.get("name") for a in t.get("annotations", [])}
        if is_auto_table(tname) or "__PBI_LocalDateTable" in annotations:
            continue
        cols = [c for c in t.get("columns", [])
                if c.get("type") != "rowNumber" and not is_auto_column(c.get("name"))]
        measures = t.get("measures", [])
        calc_cols = 0
        for c in cols:
            is_calc = c.get("type") == "calculated"
            calc_cols += is_calc
            out["Columns"].append([tname, c.get("displayFolder", ""), c.get("name", ""),
                                   c.get("dataType", ""), is_calc,
                                   join_expr(c.get("expression")) if is_calc else "",
                                   bool(c.get("isHidden", False))])
        for m in measures:
            out["Measures"].append([m.get("name", ""), join_expr(m.get("expression")),
                                    tname, m.get("displayFolder", "")])
        ttype, mode, has_m, calc_dax = "Imported", "", False, ""
        for p in t.get("partitions", []):
            src = p.get("source", {})
            mode = p.get("mode", mode) or mode
            st = src.get("type", "")
            if st == "m":
                has_m = True
                out["M Queries"].append([tname, join_expr(src.get("expression"))])
            elif st == "calculated":
                ttype, calc_dax = "Calculated Table", join_expr(src.get("expression"))
            elif st == "entity":
                ttype = "Entity (dataflow/composite)"
        if mode == "directQuery":
            ttype = "DirectQuery"
        out["Tables"].append([tname, ttype, bool(t.get("isHidden", False)), len(cols),
                              calc_cols, len(measures), mode or "import", has_m, calc_dax])
    for e in model.get("expressions", []):
        out["M Queries"].append([f"{e.get('name', '')} (shared expression / parameter)",
                                 join_expr(e.get("expression"))])
    for r in model.get("relationships", []):
        ft, tt = r.get("fromTable", ""), r.get("toTable", "")
        if is_auto_table(ft) or is_auto_table(tt):
            continue
        out["Relationships"].append([
            ft, r.get("fromColumn", ""), tt, r.get("toColumn", ""),
            f"{r.get('fromCardinality', 'many')}:{r.get('toCardinality', 'one')}",
            r.get("crossFilteringBehavior", "oneDirection"), bool(r.get("isActive", True))])


# ----------------------------------------------------------------------------
# Model extraction B: .pbix binary DataModel through pbixray
#   Hidden flag / column folders live in the model's internal SQLite metadata.
#   We capture that DB while pbixray decodes the model (several strategies).
# ----------------------------------------------------------------------------
SQLITE_MAGIC = b"SQLite format 3"


def _is_sqlite_file(p):
    try:
        with open(p, "rb") as f:
            return f.read(15) == SQLITE_MAGIC
    except Exception:
        return False


class _SqliteCapture:
    """Records every sqlite connection / temp file created while pbixray runs."""

    def __init__(self):
        self.conns, self.copies, self.diag = [], [], []
        self._keep = tempfile.mkdtemp(prefix="pbix_meta_")

    def _snapshot_tmp(self):
        found = set()
        base = tempfile.gettempdir()
        try:
            for root, dirs, files in os.walk(base):
                if root.startswith(self._keep) or root[len(base):].count(os.sep) > 2:
                    dirs[:] = []
                    continue
                for fn in files:
                    found.add(os.path.join(root, fn))
        except Exception:
            pass
        return found

    def _copy(self, path):
        try:
            if os.path.exists(path) and os.path.getsize(path) > 0 and _is_sqlite_file(path):
                dst = os.path.join(self._keep, f"{len(self.copies)}.sqlite")
                shutil.copy(path, dst)
                self.copies.append(dst)
        except Exception:
            pass

    def __enter__(self):
        self._orig = sqlite3.connect
        cap = self

        def patched(*a, **k):
            c = cap._orig(*a, **k)
            cap.conns.append(c)
            if a and isinstance(a[0], (str, os.PathLike)):
                cap._copy(str(a[0]))
            return c

        sqlite3.connect = patched
        self._before = self._snapshot_tmp()
        return self

    def __exit__(self, *exc):
        sqlite3.connect = self._orig
        for p in self._snapshot_tmp() - self._before:
            self._copy(p)
        self.diag.append(f"connections captured={len(self.conns)}, sqlite files captured={len(self.copies)}")

    def candidates(self, model):
        for c in self.conns:
            yield c
        for p in self.copies:
            try:
                yield sqlite3.connect(p)
            except Exception:
                pass
        try:
            c = _find_sqlite(model)
            if c is not None:
                yield c
        except Exception:
            pass


def _find_sqlite(obj, depth=0, seen=None):
    seen = seen if seen is not None else set()
    if id(obj) in seen or depth > 7:
        return None
    seen.add(id(obj))
    if isinstance(obj, sqlite3.Connection):
        return obj
    if isinstance(obj, (bytes, bytearray)):
        if bytes(obj[:15]) == SQLITE_MAGIC:
            fd, path = tempfile.mkstemp(suffix=".sqlite")
            with os.fdopen(fd, "wb") as f:
                f.write(obj)
            return sqlite3.connect(path)
        return None
    if isinstance(obj, str):
        if obj.lower().endswith((".sqlite", ".sqlitedb", ".db")) and os.path.exists(obj):
            return sqlite3.connect(obj)
        return None
    kids = []
    if isinstance(obj, dict):
        kids = list(obj.values())
    elif isinstance(obj, (list, tuple, set)):
        kids = list(obj)
    else:
        if hasattr(obj, "__dict__"):
            kids += list(vars(obj).values())
        for cls in type(obj).__mro__:
            for s in getattr(cls, "__slots__", ()) or ():
                try:
                    kids.append(getattr(obj, s))
                except Exception:
                    pass
    for ch in kids:
        r = _find_sqlite(ch, depth + 1, seen)
        if r is not None:
            return r
    return None


def _read_meta(conn):
    def cols_of(table):
        return {r[1].lower(): r[1] for r in conn.execute(f"PRAGMA table_info([{table}])")}

    tabs = {r[0].lower(): r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    if "column" not in tabs or "table" not in tabs:
        raise RuntimeError("no Column/Table in this DB")
    cc, tc = cols_of(tabs["column"]), cols_of(tabs["table"])
    name_parts = [f"c.{cc[k]}" for k in ("explicitname", "inferredname", "name") if k in cc]
    if not name_parts or "tableid" not in cc:
        raise RuntimeError("unexpected Column schema")
    name_expr = "COALESCE(" + ",".join(name_parts) + ")"
    hid = f"c.{cc['ishidden']}" if "ishidden" in cc else "0"
    fol = f"c.{cc['displayfolder']}" if "displayfolder" in cc else "''"
    col_info, tbl_info = {}, {}
    for t, c, h, f in conn.execute(
            f"SELECT t.{tc['name']}, {name_expr}, {hid}, {fol} FROM [{tabs['column']}] c "
            f"JOIN [{tabs['table']}] t ON t.{tc.get('id', 'ID')} = c.{cc['tableid']}"):
        col_info[(t, c)] = (bool(h), f or "")
    if "ishidden" in tc:
        for t, h in conn.execute(f"SELECT {tc['name']}, {tc['ishidden']} FROM [{tabs['table']}]"):
            tbl_info[t] = bool(h)
    if not col_info:
        raise RuntimeError("empty Column table")
    return col_info, tbl_info



def _meta_from_tmschema(m):
    """Read Is Hidden / Display Folder from pbixray's tmschema_* tables (preferred route)."""
    def frame(*names):
        attrs = [a for a in dir(m) if a.lower().startswith("tmschema_")]
        for want in names:
            for a in attrs:
                if a.lower() == want:
                    v = getattr(m, a)
                    v = v() if callable(v) else v
                    return v.to_dict("records")
        raise RuntimeError("missing " + "/".join(names) + " (have: " + ", ".join(attrs[:60]) + ")")

    def g(rec, *keys):
        low = {str(k).lower(): v for k, v in rec.items()}
        for k in keys:
            if k in low and clean(low[k]) != "":
                return low[k]
        return None

    def truthy(v):
        v = clean(v)
        return v in (True, 1, "1", "True", "true")

    tabs = frame("tmschema_tables", "tmschema_table")
    cols = frame("tmschema_columns", "tmschema_column")
    id2name, tbl_info = {}, {}
    for r in tabs:
        tid, nm = g(r, "id"), g(r, "name")
        if tid is not None and nm is not None:
            id2name[int(tid)] = str(nm)
            tbl_info[str(nm)] = truthy(g(r, "ishidden"))
    col_info = {}
    for r in cols:
        tid = g(r, "tableid")
        nm = g(r, "explicitname", "inferredname", "name")
        if tid is None or nm is None or int(tid) not in id2name:
            continue
        col_info[(id2name[int(tid)], str(nm))] = (truthy(g(r, "ishidden")), str(clean(g(r, "displayfolder")) or ""))
    if not col_info:
        raise RuntimeError("tmschema_columns was empty")
    return col_info, tbl_info


def from_pbix_binary(path, out, warnings):
    try:
        from pbixray import PBIXRay
    except ImportError:
        raise RuntimeError("The 'pbixray' package is required for .pbix files. Run: pip install pbixray")

    cap = _SqliteCapture()
    with cap:
        m = PBIXRay(path)
        frames = {}
        for attr in ("schema", "dax_columns", "dax_measures", "relationships",
                     "power_query", "m_parameters", "dax_tables"):
            try:
                frames[attr] = getattr(m, attr)
            except Exception:
                frames[attr] = None
        try:
            table_names = [str(x) for x in m.tables]
        except Exception:
            table_names = []

    col_info, tbl_info, errs = {}, {}, []
    try:
        col_info, tbl_info = _meta_from_tmschema(m)
    except Exception as e:  # noqa
        errs.append("tmschema: " + str(e))
    for cand in ([] if col_info else cap.candidates(m)):
        try:
            col_info, tbl_info = _read_meta(cand)
            break
        except Exception as e:  # noqa
            errs.append(str(e))
    known = bool(col_info)
    if not known:
        attrs = sorted(k for k in dir(m) if not k.startswith("__"))[:40]
        warnings.append(
            "Could not read 'Is Hidden' / column Folder from this .pbix (" + "; ".join(cap.diag + errs[:2]) +
            "). Everything else is complete. Please send this line to the tool author: "
            f"pbixray attrs = {attrs}. Workaround: File > Export > Power BI template (.pbit) and drop that file.")

    def recs(name):
        df = frames.get(name)
        try:
            return df.to_dict("records") if df is not None else []
        except Exception:
            return []

    dtypes = {(r["TableName"], r["ColumnName"]): clean(r.get("PandasDataType")) for r in recs("schema")}
    dax_cols = {(r["TableName"], r["ColumnName"]): clean(r.get("Expression")) for r in recs("dax_columns")}
    keys = list(dict.fromkeys(list(dtypes) + list(dax_cols)))
    n_cols, n_calc, n_meas = {}, {}, {}
    for (t, c) in sorted(keys, key=lambda k: (str(k[0]).lower(), str(k[1]).lower())):
        if is_auto_table(t) or is_auto_column(c):
            continue
        is_calc = (t, c) in dax_cols
        hidden, folder = col_info.get((t, c), ("" if not known else False, ""))
        out["Columns"].append([t, folder, c, dtypes.get((t, c), ""), is_calc,
                               dax_cols.get((t, c), "") if is_calc else "", hidden])
        n_cols[t] = n_cols.get(t, 0) + 1
        n_calc[t] = n_calc.get(t, 0) + int(is_calc)
    for r in recs("dax_measures"):
        t = r.get("TableName", "")
        if is_auto_table(t):
            continue
        out["Measures"].append([clean(r.get("Name")), clean(r.get("Expression")), t,
                                clean(r.get("DisplayFolder"))])
        n_meas[t] = n_meas.get(t, 0) + 1
    for r in recs("relationships"):
        ft, tt = r.get("FromTableName", ""), r.get("ToTableName", "")
        if is_auto_table(ft) or is_auto_table(tt):
            continue
        act = r.get("IsActive")
        out["Relationships"].append([ft, clean(r.get("FromColumnName")), tt, clean(r.get("ToColumnName")),
                                     clean(r.get("Cardinality")), clean(r.get("CrossFilteringBehavior")),
                                     True if act is None or act == "" else bool(clean(act))])
    m_tables = set()
    for r in recs("power_query"):
        m_tables.add(r.get("TableName", ""))
        out["M Queries"].append([r.get("TableName", ""), clean(r.get("Expression"))])
    for r in recs("m_parameters"):
        name = clean(r.get("ParameterName"))
        out["M Queries"].append([f"{name} (parameter)",
                                 f"{clean(r.get('Description'))}\nValue: {clean(r.get('Expression'))}".strip()])
    calc_tables = {r["TableName"]: clean(r.get("Expression")) for r in recs("dax_tables")}
    names = set(n_cols) | set(n_meas) | set(calc_tables) | m_tables | set(table_names)
    for t in sorted((n for n in names if n and not is_auto_table(n)), key=str.lower):
        ttype = ("Calculated Table" if t in calc_tables else "Imported (Power Query)" if t in m_tables
                 else "Measures only" if n_meas.get(t) and not n_cols.get(t) else "Other / Unknown")
        out["Tables"].append([t, ttype, tbl_info.get(t, "") if known else "", n_cols.get(t, 0),
                              n_calc.get(t, 0), n_meas.get(t, 0), "", t in m_tables, calc_tables.get(t, "")])


# ----------------------------------------------------------------------------
# Report layout -> visuals (report / page / visual / fields + role)
# ----------------------------------------------------------------------------
def _entity_of(expr, alias):
    sr = (expr or {}).get("SourceRef", {}) if isinstance(expr, dict) else {}
    return sr.get("Entity") or alias.get(sr.get("Source"), sr.get("Source", ""))


def _find_key(node, key):
    if isinstance(node, dict):
        if key in node:
            return node[key]
        for v in node.values():
            r = _find_key(v, key)
            if r is not None:
                return r
    elif isinstance(node, list):
        for v in node:
            r = _find_key(v, key)
            if r is not None:
                return r
    return None


def _parse_field(node, alias):
    if not isinstance(node, dict):
        return None
    if "Aggregation" in node:
        inner = _parse_field(node["Aggregation"].get("Expression"), alias)
        if inner:
            inner["agg"] = AGG.get(node["Aggregation"].get("Function"), "")
            return inner
    for kind in ("Column", "Measure"):
        b = node.get(kind)
        if isinstance(b, dict) and "Property" in b:
            return {"kind": kind, "table": _entity_of(b.get("Expression"), alias), "name": b["Property"], "agg": ""}
    if "HierarchyLevel" in node:
        hl = node["HierarchyLevel"]
        pv = _find_key(hl, "PropertyVariationSource")
        if pv:
            return {"kind": "Column", "table": _entity_of(pv.get("Expression"), alias),
                    "name": pv.get("Property", ""), "agg": f"Date hierarchy: {hl.get('Level', '')}"}
        h = hl.get("Expression", {}).get("Hierarchy", {})
        return {"kind": "Hierarchy", "table": _entity_of(h.get("Expression"), alias),
                "name": f"{h.get('Hierarchy', '')}.{hl.get('Level', '')}", "agg": ""}
    return None


def _walk_fields(node, alias):
    f = _parse_field(node, alias)
    if f:
        return [f]
    res = []
    if isinstance(node, dict):
        for v in node.values():
            res += _walk_fields(v, alias)
    elif isinstance(node, list):
        for v in node:
            res += _walk_fields(v, alias)
    return res


def _literal(d, *path):
    for p in path:
        if isinstance(d, list):
            d = d[0] if d else {}
        d = d.get(p, {}) if isinstance(d, dict) else {}
    return d.strip("'") if isinstance(d, str) else ""


def _textbox_text(obj):
    try:
        paras = obj["general"][0]["properties"]["paragraphs"]
        txt = " ".join(r.get("value", "") for p in paras for r in p.get("textRuns", []))
        return txt.strip()
    except Exception:
        return ""


def _add_visual(out, page, vtype, title, vid, pos, fields):
    seen, uniq = set(), []
    for f in fields:
        k = (f["kind"], f["table"], f["name"], f.get("role", ""), f.get("agg", ""))
        if k not in seen:
            seen.add(k)
            uniq.append(f)
    out["_visuals"].append({"page": page, "vtype": vtype, "title": title, "vid": vid,
                            "pos": pos, "fields": uniq})


def _json(v):
    if isinstance(v, str):
        try:
            return json.loads(v)
        except Exception:
            return None
    return v


def parse_layout(zf, out, warnings):
    names = zf.namelist()
    lname = next((n for n in names if n.lower() == "report/layout"), None)
    if lname:
        layout = json.loads(decode_bytes(zf.read(lname)))
        # report-level measures
        cfg0 = _json(layout.get("config")) or {}
        for ext in cfg0.get("modelExtensions", []) or []:
            for ent in ext.get("entities", []):
                for ms in ent.get("measures", []):
                    out["Measures"].append([ms.get("name", ""), ms.get("expression", ""),
                                            ent.get("name", ""), "(report-level measure)"])
        rf = []
        for f in _json(layout.get("filters")) or []:
            for fld in _walk_fields(f.get("expression"), {}):
                fld["role"] = "Report Filter"
                rf.append(fld)
        if rf:
            _add_visual(out, "(All pages)", "Report Filters", "", "", (0, 0, 0, 0), rf)
        for sec in layout.get("sections", []):
            page = sec.get("displayName") or sec.get("name", "")
            pf = []
            for f in _json(sec.get("filters")) or []:
                for fld in _walk_fields(f.get("expression"), {}):
                    fld["role"] = "Page Filter"
                    pf.append(fld)
            if pf:
                _add_visual(out, page, "Page Filters", "", "", (0, 0, 0, 0), pf)
            for vc in sec.get("visualContainers", []):
                cfg = _json(vc.get("config")) or {}
                sv = cfg.get("singleVisual")
                if not sv:
                    continue
                pq = sv.get("prototypeQuery") or {}
                alias = {f.get("Name"): f.get("Entity") for f in pq.get("From", [])}
                roles = {}
                for role, items in (sv.get("projections") or {}).items():
                    for it in items:
                        roles.setdefault(it.get("queryRef"), []).append(role)
                fields = []
                for sel in pq.get("Select", []):
                    for fld in _walk_fields(sel, alias):
                        fld["role"] = ", ".join(roles.get(sel.get("Name"), [])) or "Other (sort / tooltip / condition)"
                        fields.append(fld)
                for f in _json(vc.get("filters")) or []:
                    for fld in _walk_fields(f.get("expression"), alias):
                        fld["role"] = "Visual Filter"
                        fields.append(fld)
                title = _literal(sv.get("vcObjects", {}), "title", "properties", "text", "expr", "Literal", "Value")
                if not title and sv.get("visualType") == "textbox":
                    title = _textbox_text(sv.get("objects", {}))
                _add_visual(out, page, sv.get("visualType", ""), title, cfg.get("name", ""),
                            (round(vc.get("x", 0)), round(vc.get("y", 0)),
                             round(vc.get("width", 0)), round(vc.get("height", 0))), fields)
        return
    # PBIR folder format
    vis_files = [n for n in names if re.search(r"report/definition/pages/[^/]+/visuals/[^/]+/visual\.json$", n, re.I)]
    if not vis_files:
        warnings.append("No report layout found - Visual Design / Lineage tabs will be empty.")
        return
    pnames = {}
    for n in names:
        mm = re.search(r"pages/([^/]+)/page\.json$", n, re.I)
        if mm:
            try:
                pnames[mm.group(1)] = json.loads(decode_bytes(zf.read(n))).get("displayName", mm.group(1))
            except Exception:
                pass
    for n in vis_files:
        pid = re.search(r"pages/([^/]+)/visuals", n, re.I).group(1)
        v = json.loads(decode_bytes(zf.read(n)))
        vis = v.get("visual", {})
        fields = []
        for role, st in (vis.get("query", {}).get("queryState", {}) or {}).items():
            for pr in st.get("projections", []):
                for fld in _walk_fields(pr.get("field"), {}):
                    fld["role"] = role
                    fields.append(fld)
        for f in (v.get("filterConfig", {}) or {}).get("filters", []):
            for fld in _walk_fields(f.get("field"), {}):
                fld["role"] = "Visual Filter"
                fields.append(fld)
        pos = v.get("position", {})
        title = _literal(vis.get("visualContainerObjects", {}), "title", "properties", "text", "expr", "Literal", "Value")
        _add_visual(out, pnames.get(pid, pid), vis.get("visualType", ""), title, v.get("name", ""),
                    (round(pos.get("x", 0)), round(pos.get("y", 0)),
                     round(pos.get("width", 0)), round(pos.get("height", 0))), fields)


# ----------------------------------------------------------------------------
# M parsing: source system / object / steps
# ----------------------------------------------------------------------------
NON_SOURCE = {"Table", "List", "Text", "Number", "Date", "DateTime", "DateTimeZone", "Duration", "Record",
              "Binary", "Logical", "Time", "Currency", "Percentage", "Int64", "Double", "Single", "Decimal",
              "Function", "Error", "Expression", "Uri", "Splitter", "Combiner", "Replacer", "Comparer",
              "Order", "Occurrence", "JoinKind", "Character", "Culture", "Diagnostics", "Variable",
              "Type", "Any", "Action", "RelativePosition", "Precision", "ExtraValues", "MissingField",
              "GroupKind", "TraceLevel", "Day", "Value", "Lines", "Json", "Xml", "Csv", "Excel", "Parquet",
              "Pdf", "Access", "Embedded", "Graph", "Progress", "Module", "Cube", "Fabric"}
FILE_FORMATS = {"Csv.Document", "Excel.Workbook", "Json.Document", "Xml.Tables", "Xml.Document",
                "Parquet.Document", "Pdf.Tables", "Lines.FromBinary", "Access.Database", "Excel.CurrentWorkbook"}
DB2 = {"Sql", "PostgreSQL", "MySQL", "AmazonRedshift", "Sql.Databases"}
STR = r'"((?:[^"]|"")*)"'
CALL_RE = re.compile(r"\b([A-Z][A-Za-z0-9]*)\.([A-Z][A-Za-z0-9]*)\s*\(")


def _unq(s):
    return (s or "").replace('""', '"').replace("#(lf)", "\n").replace("#(cr)", "").replace("#(tab)", "\t")


def m_steps(m):
    lines = m.splitlines()
    steps, indent = [], None
    for ln in lines:
        mm = re.match(r'^(\s*)(#"(?:[^"]|"")+"|[A-Za-z_]\w*)\s*=\s*(.*)$', ln)
        if not mm:
            continue
        if indent is None:
            indent = len(mm.group(1))
        if len(mm.group(1)) != indent:
            continue
        name = mm.group(2)
        name = name[2:-1] if name.startswith('#"') else name
        if name.lower() == "let":
            continue
        fn = CALL_RE.search(mm.group(3))
        steps.append(f"{len(steps) + 1}. {name}" + (f"  [{fn.group(1)}.{fn.group(2)}]" if fn else ""))
    return "\n".join(steps)


def parse_source(m, queries, self_name="", depth=0):
    res = {"system": "", "server": "", "object": "", "sql": ""}
    if not m:
        return res
    call = None
    for mm in CALL_RE.finditer(m):
        pre, fn = mm.group(1), mm.group(1) + "." + mm.group(2)
        if fn == "Value.NativeQuery" or fn in FILE_FORMATS or pre not in NON_SOURCE or fn in (
                "File.Contents", "Web.Contents"):
            if fn in ("Value.NativeQuery",):
                continue
            call = mm
            break
    if call:
        fn = call.group(1) + "." + call.group(2)
        pre = call.group(1)
        tail = m[call.end():call.end() + 1500]
        if fn in FILE_FORMATS:
            inner = re.search(r"(File|Web)\.Contents\s*\(\s*" + STR, m) or \
                re.search(r"(SharePoint\.\w+|Folder\.\w+)\s*\(\s*" + STR, m)
            res["system"] = f"{fn.split('.')[0]} file" if inner else fn.split(".")[0]
            res["server"] = _unq(inner.group(2)) if inner else ""
        elif fn in ("File.Contents", "Web.Contents"):
            lit = re.match(r"\s*" + STR, tail)
            res["system"] = "Web" if fn.startswith("Web") else "File"
            res["server"] = _unq(lit.group(1)) if lit else ""
        else:
            lits = re.match(r"\s*" + STR + r"(?:\s*,\s*" + STR + ")?", tail)
            res["system"] = pre if fn.endswith((".Databases", ".Database")) or True else fn
            if lits:
                res["server"] = _unq(lits.group(1))
                if lits.group(2):
                    res["server"] += f"  ({'database' if pre in DB2 else ('warehouse' if pre == 'Snowflake' else 'db')}: {_unq(lits.group(2))})"
                    if pre in DB2:
                        res["object"] = _unq(lits.group(2))
        parts = []
        for rec in re.finditer(r"\[\s*((?:\w+\s*=\s*" + STR + r"\s*,?\s*)+)\]", m):
            kv = dict((a, _unq(b)) for a, b in re.findall(r"(\w+)\s*=\s*" + STR, rec.group(1)))
            if "Kind" in kv and kv.get("Kind") in ("Database", "Schema", "Table", "View", "Sheet", "Folder", "Catalog"):
                parts.append(kv.get("Name") or kv.get("Item", ""))
            elif "Schema" in kv and "Item" in kv:
                parts += [kv["Schema"], kv["Item"]]
            elif "Item" in kv and "Kind" not in kv:
                parts.append(kv["Item"])
        if parts:
            res["object"] = ".".join(x for x in ([res["object"]] if res["object"] and res["object"] not in parts else []) + parts if x)
        nq = re.search(r"Value\.NativeQuery\s*\(.*?,\s*" + STR, m, re.S) or re.search(r"\bQuery\s*=\s*" + STR, m, re.S)
        if nq:
            res["sql"] = _unq(nq.group(1))[:1500]
        return res
    if re.search(r"#table\s*\(|Table\.FromRows\s*\(|Table\.FromColumns\s*\(", m):
        res["system"] = "Manual entry (Enter Data)"
        return res
    if depth < 5:
        for qn, qm in queries.items():
            if qn == self_name or "(" in qn:
                continue
            if re.search(r'(?<![\w.])(#"%s"|%s)(?![\w])' % (re.escape(qn), re.escape(qn)), m):
                sub = parse_source(qm, queries, qn, depth + 1)
                if sub["system"]:
                    sub["server"] = sub["server"]
                    return sub
    return res


# ----------------------------------------------------------------------------
# M column tracing: model column -> original source column (through renames etc.)
# ----------------------------------------------------------------------------
def _scan_str(s, i):
    j, n = i + 1, len(s)
    while j < n:
        if s[j] == '"':
            if j + 1 < n and s[j + 1] == '"':
                j += 2
                continue
            return j + 1
        j += 1
    return n


def _strip_comments(s):
    out, i, n = [], 0, len(s)
    while i < n:
        if s[i] == '"':
            j = _scan_str(s, i)
            out.append(s[i:j])
            i = j
        elif s.startswith("//", i):
            j = s.find("\n", i)
            i = n if j < 0 else j
        elif s.startswith("/*", i):
            j = s.find("*/", i)
            i = n if j < 0 else j + 2
        else:
            out.append(s[i])
            i += 1
    return "".join(out)


_SPLIT_CACHE = {}


def _m_split_let(m):
    if m in _SPLIT_CACHE:
        return _SPLIT_CACHE[m]
    s = _strip_comments(m).strip()
    if not re.match(r"let\b", s, re.I):
        res = ({"__expr__": s}, "__expr__")
        _SPLIT_CACHE[m] = res
        return res
    body, n, i, depth, start = s[3:], len(s) - 3, 0, 0, 0
    stmts, final = [], None
    broke = False
    while i < n:
        ch = body[i]
        if ch == '"':
            i = _scan_str(body, i)
            continue
        if ch in "([{":
            depth += 1
        elif ch in ")]}":
            depth -= 1
        elif depth == 0:
            if ch == ",":
                stmts.append(body[start:i])
                start = i + 1
            elif ch == "i" and body.startswith("in", i) and (i == 0 or body[i - 1].isspace()) \
                    and (i + 2 >= n or body[i + 2].isspace()):
                stmts.append(body[start:i])
                final = body[i + 2:].strip()
                broke = True
                break
        i += 1
    if not broke:
        stmts.append(body[start:])
    d = {}
    for st in stmts:
        mm = re.match(r'\s*(#"(?:[^"]|"")*"|[A-Za-z_]\w*)\s*=\s*(.*)$', st, re.S)
        if mm:
            nm = mm.group(1)
            nm = _unq(nm[2:-1]) if nm.startswith('#"') else nm
            d[nm] = mm.group(2).strip()
    if final:
        final = _unq(final[2:-1]) if final.startswith('#"') and final.endswith('"') else final.split()[0]
    res = (d, final or (list(d)[-1] if d else ""))
    _SPLIT_CACHE[m] = res
    return res


def _call_args(expr):
    m = re.match(r"\s*([A-Z]\w*(?:\.\w+)+)\s*\(", expr or "")
    if not m:
        return None, []
    i, n, depth, start, args = m.end(), len(expr), 1, m.end(), []
    while i < n:
        ch = expr[i]
        if ch == '"':
            i = _scan_str(expr, i)
            continue
        if ch in "([{":
            depth += 1
        elif ch in ")]}":
            depth -= 1
            if depth == 0:
                args.append(expr[start:i])
                break
        elif ch == "," and depth == 1:
            args.append(expr[start:i])
            start = i + 1
        i += 1
    return m.group(1), [a.strip() for a in args]


def _lits(s):
    return [_unq(x) for x in re.findall(r'"((?:[^"]|"")*)"', s or "")]


def _ref(a):
    mm = re.match(r'\s*(#"((?:[^"]|"")*)"|[A-Za-z_]\w*)', a or "")
    if not mm:
        return None
    return _unq(mm.group(2)) if mm.group(1).startswith('#"') else mm.group(1)


def _refs_in_expr(expr):
    names = []
    for mm in re.finditer(r'\[\s*(?:#"((?:[^"]|"")+)"|([^\]\[=",]+?))\s*\]', expr or ""):
        nm = _unq(mm.group(1)) if mm.group(1) else mm.group(2).strip()
        if nm and nm not in names:
            names.append(nm)
    return names


def _sql_map(sql, cur):
    s = re.sub(r"--[^\n]*|/\*.*?\*/", " ", sql, flags=re.S)
    m = re.search(r"\bselect\b", s, re.I)
    if not m:
        return cur, ""
    i, n, depth, start, items, frm = m.end(), len(s), 0, m.end(), [], ""
    while i < n:
        ch = s[i]
        if ch == "'":
            j = s.find("'", i + 1)
            i = j + 1 if j >= 0 else n
            continue
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth -= 1
        elif depth == 0:
            if ch == ",":
                items.append(s[start:i])
                start = i + 1
            elif s[i:i + 4].lower() == "from" and s[i - 1].isspace() and i + 4 < n and s[i + 4].isspace():
                items.append(s[start:i])
                fm = re.match(r'\s+([\w."\[\]`$]+)', s[i + 4:])
                frm = fm.group(1) if fm else ""
                break
        i += 1
    for it in items:
        it = it.strip()
        am = re.search(r'\s+as\s+("[^"]+"|\[[^\]]+\]|\w+)\s*$', it, re.I)
        if am:
            alias, expr = am.group(1).strip('"[]'), it[:am.start()].strip()
        else:
            expr, alias = it, it.split(".")[-1].strip('"[] ')
        if alias.lower() == cur.lower():
            return expr, frm
    return cur, frm


def _src_table_col(p, cur):
    obj, sql, col = p.get("object", ""), p.get("sql", ""), cur
    if sql:
        col, frm = _sql_map(sql, cur) if cur else (cur, "")
        tbl = frm.split(".")[-1] if frm else obj
    elif obj:
        tbl = obj.split(".")[-1] if "file" not in p.get("system", "").lower() else obj
    else:
        tbl = os.path.basename((p.get("server", "") or "").replace("\\", "/").split("  (")[0])
    return (tbl or "").strip('"[]` '), col


def _src_result(queries, qname, cur, notes):
    p = parse_source(queries[qname], queries, qname)
    tbl, col = _src_table_col(p, cur)
    return {"src_table": tbl, "src_col": col, "system": p["system"], "server": p["server"],
            "object": p["object"], "sql": p["sql"], "notes": list(notes), "type": "Power Query (M)"}


def trace_m_column(queries, qname, col, depth=0):
    """Walk the M steps of query `qname` backwards from final output column `col`."""
    if depth > 6 or qname not in queries:
        return []
    stmts, final = _m_split_let(queries[qname])
    nested = {}
    for nm, ex in stmts.items():
        fn, args = _call_args(ex)
        if fn == "Table.NestedJoin" and len(args) > 4:
            L = _lits(args[4])
            other = _ref(args[2])
            if other in stmts and other not in queries:
                other = _ref(stmts[other]) or other
            if L:
                nested[L[0]] = other
    results = []

    def lc(a, b):
        return a.lower() == b.lower()

    def sub(qn, c, notes):
        for r in trace_m_column(queries, qn, c, depth + 1):
            results.append(dict(r, notes=notes + r["notes"]))

    def walk(step, cur, notes, hops):
        if hops > 300:
            return
        ex = stmts.get(step)
        if ex is None:
            if step in queries and step != qname:
                sub(step, cur, notes)
            return
        fn, args = _call_args(ex)
        if fn is None:
            ref = _ref(ex)
            if ref and ref != step and (ref in stmts or ref in queries):
                return walk(ref, cur, notes, hops + 1)
            results.append(_src_result(queries, qname, cur, notes))
            return
        prefix = fn.split(".")[0]
        if not fn.startswith("Table.") or fn in ("Table.FromRows", "Table.FromColumns", "Table.FromList",
                                                 "Table.FromRecords"):
            results.append(_src_result(queries, qname, cur, notes))
            return
        prev = _ref(args[0]) if args else None
        if not prev:
            results.append(_src_result(queries, qname, cur, notes))
            return

        def cont(c=cur, n=notes):
            return walk(prev, c, n, hops + 1)

        if fn == "Table.RenameColumns" and len(args) > 1:
            L = _lits(args[1])
            for k in range(0, len(L) - 1, 2):
                if lc(L[k + 1], cur):
                    return walk(prev, L[k], notes + [f'Renamed "{L[k]}" -> "{L[k + 1]}" (step {step})'], hops + 1)
            return cont()
        if fn == "Table.DuplicateColumn" and len(args) > 2:
            L1, L2 = _lits(args[1]), _lits(args[2])
            if L1 and L2 and lc(L2[0], cur):
                return walk(prev, L1[0], notes + [f'Duplicated "{L1[0]}" as "{L2[0]}" (step {step})'], hops + 1)
            return cont()
        if fn.startswith("Table.Add") and len(args) > 1:
            L = _lits(args[1])
            if L and lc(L[0], cur):
                body = args[2] if len(args) > 2 else ""
                refs = _refs_in_expr(body)
                note = f'Custom column "{cur}" created in Power Query (step {step})' + \
                       (f": {' '.join(body.split())[:220]}" if body else "")
                if refs:
                    for r in refs:
                        walk(prev, r, notes + [note], hops + 1)
                    return
                r = _src_result(queries, qname, "", notes + [note])
                r["src_col"] = ""
                results.append(r)
                return
            return cont()
        if fn in ("Table.ExpandTableColumn", "Table.ExpandRecordColumn") and len(args) > 3:
            orig = _lits(args[2])
            new = _lits(args[3]) or orig
            for k, nm in enumerate(new):
                if lc(nm, cur) and k < len(orig):
                    nc = (_lits(args[1]) or [""])[0]
                    other = nested.get(nc)
                    n2 = notes + [f'Expanded "{orig[k]}" from joined column "{nc}"' +
                                  (f" of query {other}" if other else "") +
                                  (f' as "{nm}"' if nm != orig[k] else "") + f" (step {step})"]
                    if other and other in queries:
                        sub(other, orig[k], n2)
                        return
                    return walk(prev, orig[k], n2, hops + 1)
            return cont()
        if fn == "Table.CombineColumns" and len(args) > 3:
            L = _lits(args[3])
            if L and lc(L[0], cur):
                for r in _lits(args[1]):
                    walk(prev, r, notes + [f'Combined columns {_lits(args[1])} into "{cur}" (step {step})'], hops + 1)
                return
            return cont()
        if fn == "Table.SplitColumn" and len(args) > 1:
            src = (_lits(args[1]) or [""])[0]
            names = _lits(args[3]) if len(args) > 3 else []
            if src and ((names and any(lc(x, cur) for x in names)) or
                        re.match(re.escape(src) + r"\.\d+$", cur, re.I)):
                return walk(prev, src, notes + [f'Split from "{src}" (step {step})'], hops + 1)
            return cont()
        if fn == "Table.Group" and len(args) > 2:
            if any(lc(x, cur) for x in _lits(args[2])):
                refs = [r for r in _refs_in_expr(args[2]) if not lc(r, cur)]
                note = f'Aggregated in Group By (step {step})'
                if refs:
                    for r in refs:
                        walk(prev, r, notes + [note], hops + 1)
                    return
            return cont()
        if fn.startswith("Table.Unpivot") and len(args) > 3:
            if any(lc(x, cur) for x in _lits(args[2]) + _lits(args[3])):
                r = _src_result(queries, qname, "", notes + [f'Produced by unpivot (step {step})'])
                r["src_col"] = ""
                results.append(r)
                return
            return cont()
        if fn == "Table.PrefixColumns" and len(args) > 1:
            pf = (_lits(args[1]) or [""])[0]
            if pf and cur.lower().startswith(pf.lower() + "."):
                return walk(prev, cur[len(pf) + 1:], notes + [f'Prefix "{pf}." removed (step {step})'], hops + 1)
            return cont()
        if fn == "Table.TransformColumnNames":
            return walk(prev, cur, notes + [f"Column names changed by a function (step {step}); name assumed unchanged"], hops + 1)
        if fn == "Table.PromoteHeaders":
            return walk(prev, cur, notes + [f"Headers promoted from first row (step {step})"], hops + 1)
        if fn == "Table.Combine" and args:
            toks = re.findall(r'#"(?:[^"]|"")*"|[A-Za-z_]\w*', args[0])
            refs = [_unq(t[2:-1]) if t.startswith('#"') else t for t in toks]
            for r in refs:
                if r in stmts:
                    walk(r, cur, notes + [f"Appended queries (step {step})"], hops + 1)
                elif r in queries:
                    sub(r, cur, notes + [f"Appended query {r} (step {step})"])
            return
        return cont()

    walk(final, col, [], 0)
    seen, uniq = set(), []
    for r in results:
        k = (r["src_table"], r["src_col"], r["object"], tuple(r["notes"]))
        if k not in seen:
            seen.add(k)
            uniq.append(r)
    return uniq


# ----------------------------------------------------------------------------
# DAX dependency resolution + lineage
# ----------------------------------------------------------------------------
COMMENT_RE = re.compile(r"//[^\n]*|/\*.*?\*/", re.S)
STR_RE = re.compile(r'"(?:[^"]|"")*"')
QUAL_RE = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z_]\w*))\s*\[([^\]]+)\]")
BARE_RE = re.compile(r"\[([^\]]+)\]")
TOK_RE = re.compile(r"'((?:[^']|'')+)'|\b([A-Za-z_]\w*)\b(?!\s*\()")


class Ctx:
    def __init__(self, out):
        self.columns = {(r[0], r[2]): r for r in out["Columns"]}
        self.calc_cols = {(r[0], r[2]): r[5] for r in out["Columns"] if r[4]}
        self.measures = {}
        for r in out["Measures"]:
            self.measures.setdefault(r[0], (r[2], r[1]))
        self.tables = {r[0] for r in out["Tables"]} | {r[2] for r in out["Measures"]}
        self.calc_tables = {r[0]: r[8] for r in out["Tables"] if str(r[1]).startswith("Calculated")}
        self.mq = {r[0]: r[1] for r in out["M Queries"]}
        self._src = {}

    def dax_refs(self, expr, own):
        s = STR_RE.sub('""', COMMENT_RE.sub(" ", expr or ""))
        cols, meas, tabs = [], [], []

        def repl(m):
            t = (m.group(1) or m.group(2)).replace("''", "'")
            n = m.group(3)
            if (t, n) in self.columns:
                cols.append((t, n))
            elif n in self.measures:
                meas.append(n)
            elif t in self.tables:
                cols.append((t, n))
            return " "

        s2 = QUAL_RE.sub(repl, s)
        for m in BARE_RE.finditer(s2):
            n = m.group(1)
            if n in self.measures:
                meas.append(n)
            elif (own, n) in self.columns:
                cols.append((own, n))
        for m in TOK_RE.finditer(s2):
            t = (m.group(1) or m.group(2) or "").replace("''", "'")
            if t in self.tables:
                tabs.append(t)
        return list(dict.fromkeys(cols)), list(dict.fromkeys(meas)), list(dict.fromkeys(tabs))

    def resolve_expr(self, expr, own, via, seen):
        cols, meas, tabs = self.dax_refs(expr, own)
        res = []
        for t, c in cols:
            res += self.resolve_column(t, c, via, seen)
        for n in meas:
            res += self.resolve_measure(n, via, seen)
        used = {t for t, _ in cols}
        for t in tabs:
            if t not in used:
                res.append((t, "", via))
        return res

    def resolve_measure(self, name, via, seen):
        key = ("m", name)
        if key in seen:
            return []
        m = self.measures.get(name)
        if not m:
            return [("?", name, via + [f"[{name}] (not found in model)"])]
        return self.resolve_expr(m[1], m[0], via + [f"[{name}]"], seen | {key})

    def resolve_column(self, t, c, via, seen):
        key = ("c", t, c)
        if key in seen:
            return []
        if t in self.calc_tables and (t, c) not in self.calc_cols:
            cols = [x for x in self.dax_refs(self.calc_tables[t], t)[0] if x[0] != t]
            match = [x for x in cols if x[1].lower() == c.lower()] or cols
            if match:
                res = []
                for t2, c2 in match:
                    res += self.resolve_column(t2, c2, via + [f"{t}[{c}] (column of calculated table)"], seen | {key})
                return res
        f = self.calc_cols.get((t, c))
        if f:
            return self.resolve_expr(f, t, via + [f"{t}[{c}] (calculated column)"], seen | {key})
        return [(t, c, via)]

    def table_sources(self, t, depth=0):
        if t in self._src:
            return self._src[t]
        self._src[t] = [{"type": "Unknown", "system": "", "server": "", "object": "", "sql": ""}]
        if t in self.calc_tables:
            expr = self.calc_tables[t]
            refs = [x for x in self.dax_refs(expr, t)[2] if x != t] + [c[0] for c in self.dax_refs(expr, t)[0] if c[0] != t]
            res = []
            if depth < 4:
                for r in dict.fromkeys(refs):
                    for s in self.table_sources(r, depth + 1):
                        res.append(dict(s, type="Calculated Table (DAX) <- " + s["type"]))
            self._src[t] = res or [{"type": "Calculated Table (DAX)", "system": "DAX", "server": "",
                                    "object": "(calculated)", "sql": (expr or "")[:300]}]
        elif t in self.mq:
            p = parse_source(self.mq[t], self.mq, t)
            self._src[t] = [dict(p, type="Power Query (M)")]
        elif t == "?":
            self._src[t] = [{"type": "Unresolved", "system": "", "server": "", "object": "", "sql": ""}]
        else:
            self._src[t] = [{"type": "No source (measure table / not found)", "system": "", "server": "",
                             "object": "", "sql": ""}]
        return self._src[t]


def _column_sources(ctx, t, c):
    key = ("src", t, c)
    if key in ctx._src:
        return ctx._src[key]
    res = []
    if c and t in ctx.mq:
        res = trace_m_column(ctx.mq, t, c)
    if not res:
        for s_ in ctx.table_sources(t):
            tbl, col = _src_table_col(s_, c)
            res.append({"src_table": tbl, "src_col": c if c else "", "system": s_["system"], "server": s_["server"],
                        "object": s_["object"], "sql": s_["sql"],
                        "notes": ([] if not c else ["Column name assumed unchanged (could not trace through M)"]),
                        "type": s_["type"]})
    ctx._src[key] = res
    return res


def finalize(out, report):
    ctx = Ctx(out)
    # M Queries sheet (name, expr) -> full row
    full = []
    for name, expr in [(r[0], r[1]) for r in out["M Queries"]]:
        p = parse_source(expr, ctx.mq, name)
        srv = p["server"] + (f"\nSQL: {p['sql']}" if p["sql"] else "")
        full.append([name, p["system"], p["server"], p["object"] if not p["sql"] else (p["object"] or "(native query)"),
                     m_steps(expr), expr])
        if p["sql"]:
            full[-1][2] = srv
    out["M Queries"] = full

    vd, lin = [], []
    impact = {}
    for v in out["_visuals"]:
        x, y, w, h = v["pos"]
        base = [report, v["page"], v["vtype"], v["title"]]
        vlabel = f"{v['page']} > {v['title'] or v['vtype']}" + (f" [{v['vid'][:6]}]" if v["vid"] else "")
        if not v["fields"]:
            vd.append(base + ["", "", "", "", "", x, y, w, h, v["vid"]])
            continue
        for f in v["fields"]:
            vd.append(base + [f["name"], f["kind"], f["table"], f.get("role", ""), f.get("agg", ""), x, y, w, h, v["vid"]])
            formula = ""
            if f["kind"] == "Measure":
                formula = ctx.measures.get(f["name"], ("", ""))[1]
                bases = ctx.resolve_measure(f["name"], [], set())
                flabel = f"[{f['name']}]"
            elif f["kind"] == "Column":
                formula = ctx.calc_cols.get((f["table"], f["name"]), "")
                bases = ctx.resolve_column(f["table"], f["name"], [], set())
                flabel = f"{f['table']}[{f['name']}]"
            else:
                bases = [(f["table"], f["name"], [])]
                flabel = f"{f['table']}[{f['name']}]"
            if not bases:
                bases = [(f["table"], "", [])]
            seen_rows = set()
            for bt, bc, via in bases:
                for s in _column_sources(ctx, bt, bc):
                    chain = list(via)
                    if not chain or not chain[0].startswith(flabel):
                        chain = [flabel] + chain
                    blabel = f"{bt}[{bc}]" if bc else f"{bt} (table)"
                    if not chain[-1].startswith(blabel):
                        chain.append(blabel)
                    src_lbl = f"{s['src_table']}[{s['src_col']}]" if s["src_col"] else (s["src_table"] or "(no source column)")
                    path = " -> ".join([v["page"], f"{v['vtype']}: {v['title'] or v['vid'][:8] or '(untitled)'}"] + chain +
                                       [f"SOURCE {src_lbl}" + (f" [{s['system']}: {s['object']}]" if s["system"] else "")])
                    pq = "\n".join(s["notes"])
                    row = base + [f["name"] + (f" ({f['agg']})" if f.get("agg") else ""), f["kind"], f["table"], formula,
                                  " -> ".join(via), bt, bc, pq, s["src_table"], s["src_col"], s["system"], s["server"],
                                  s["object"], path, s["sql"]]
                    k = tuple(map(str, row))
                    if k in seen_rows:
                        continue
                    seen_rows.add(k)
                    lin.append(row)
                    ik = (s["system"], s["object"], s["src_table"], s["src_col"], bt, bc)
                    e = impact.setdefault(ik, {"chains": [], "pq": [], "pages": set(), "visuals": []})
                    if " -> ".join(via) not in e["chains"] and via:
                        e["chains"].append(" -> ".join(via))
                    if pq and pq not in e["pq"]:
                        e["pq"].append(pq)
                    e["pages"].add(v["page"])
                    if vlabel not in e["visuals"]:
                        e["visuals"].append(vlabel)
    imp = []
    for (system, obj, st, sc, bt, bc), e in sorted(impact.items(), key=lambda kv: tuple(str(z).lower() for z in kv[0][2:4] + kv[0][4:])):
        imp.append([system, obj, st, sc or "(table level)", bt, bc, "\n".join(e["chains"]), "\n".join(e["pq"]),
                    ", ".join(sorted(e["pages"])), len(e["visuals"]), "\n".join(e["visuals"])])
    out["Impact Analysis"] = imp
    out["Visual Design"], out["Lineage"] = vd, lin
    out["Tables"].sort(key=lambda r: str(r[0]).lower())
    out["Measures"].sort(key=lambda r: (str(r[2]).lower(), str(r[0]).lower()))
    out["Relationships"].sort(key=lambda r: (str(r[0]).lower(), str(r[2]).lower()))


def inspect(path, report=None):
    report = report or os.path.splitext(os.path.basename(path))[0]
    out, warnings = new_result()
    with zipfile.ZipFile(path) as zf:
        names = {n.lower(): n for n in zf.namelist()}
        if "datamodelschema" in names:
            from_schema(json.loads(decode_bytes(zf.read(names["datamodelschema"]))), out)
        elif "datamodel" in names:
            from_pbix_binary(path, out, warnings)
        else:
            warnings.append("No data model in this file (live-connected report?). Only visuals could be read.")
        try:
            parse_layout(zf, out, warnings)
        except Exception as e:
            warnings.append(f"Could not parse report layout: {type(e).__name__}: {e}")
    finalize(out, report)
    return out, warnings


# ----------------------------------------------------------------------------
# Excel
# ----------------------------------------------------------------------------
SHEET_DESC = {
    "Tables": "Every table with type, hidden flag, column / measure counts, calculated-table DAX",
    "Columns": "Every column: folder, data type, calculated or not, formula, hidden or not",
    "Measures": "Every measure with formula, table and display folder",
    "Relationships": "Model relationships: columns, cardinality, filter direction, active",
    "M Queries": "Power Query code with source system, source object and transformation steps",
    "Visual Design": "Every visual: page, type, title and each field used with its role",
    "Lineage": "Visual -> column/measure -> dependent columns -> Power Query -> source column",
    "Impact Analysis": "Source column -> model column -> pages / visuals affected if it changes",
}


def summary_info(data, report, warnings):
    cols, meas = data["Columns"], data["Measures"]
    vis = [v for v in data["_visuals"] if v["vtype"] not in ("Page Filters", "Report Filters")]
    pages = {}
    for v in vis:
        e = pages.setdefault(v["page"], {"visuals": set(), "fields": set()})
        e["visuals"].add(v["vid"] or id(v))
        for f in v["fields"]:
            e["fields"].add((f["table"], f["name"]))
    systems = {}
    for r in data["M Queries"]:
        if "(parameter)" in str(r[0]) or "(shared expression" in str(r[0]):
            continue
        systems.setdefault(r[1] or "(not detected)", []).append(r[0])
    src_cols = {(r[2], r[3]) for r in data["Impact Analysis"] if r[3] and r[3] != "(table level)"}
    metrics = [
        ("Tables", len(data["Tables"])),
        ("Columns", len(cols)),
        ("   of which calculated columns", sum(1 for r in cols if r[4] is True)),
        ("   of which hidden columns", sum(1 for r in cols if r[6] is True)),
        ("Measures", len(meas)),
        ("Relationships", len(data["Relationships"])),
        ("Power Query queries", sum(len(v) for v in systems.values())),
        ("Pages", len(pages)),
        ("Visuals", sum(len(e["visuals"]) for e in pages.values())),
        ("Distinct source columns feeding visuals", len(src_cols)),
    ]
    return metrics, pages, systems


def _write_summary(wb, data, selected, report, warnings):
    from datetime import datetime
    ws = wb.create_sheet("Summary", 0)
    ws.sheet_view.showGridLines = False
    navy = PatternFill("solid", start_color="1F3864")
    light = PatternFill("solid", start_color="E8EEF9")
    metrics, pages, systems = summary_info(data, report, warnings)
    ws["A1"] = "PBIX Inspector - Summary"
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color="1F3864")
    ws["A3"], ws["B3"] = "Report", report
    ws["A4"], ws["B4"] = "Generated", datetime.now().strftime("%Y-%m-%d %H:%M")
    ws["A5"], ws["B5"] = "Sheets included", f"{len(selected)} of {len(SHEETS)} (plus this summary)"
    for r in (3, 4, 5):
        ws.cell(r, 1).font = Font(name="Arial", bold=True, size=10)
        ws.cell(r, 2).font = Font(name="Arial", size=10)
    row = 7

    def section(title, headers):
        nonlocal row
        ws.cell(row, 1, title).font = Font(name="Arial", bold=True, size=12, color="1F3864")
        row += 1
        for i, h in enumerate(headers, 1):
            c = ws.cell(row, i, h)
            c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
            c.fill = navy
        row += 1

    def put(vals, link=None):
        nonlocal row
        for i, v in enumerate(vals, 1):
            c = ws.cell(row, i, v)
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if link and i == 1:
                c.hyperlink = f"#'{link}'!A1"
                c.font = Font(name="Arial", size=10, color="0563C1", underline="single")
        row += 1

    section("Contents (sheets in this file)", ["Sheet", "What it contains", "Rows"])
    for n in selected:
        put([n, SHEET_DESC.get(n, ""), len(data[n])], link=n)
    skipped = [n for n in SHEETS if n not in selected]
    if skipped:
        put(["Not included", ", ".join(skipped), ""])
    row += 1
    section("Key metrics", ["Metric", "Value"])
    for k, v in metrics:
        put([k, v])
    row += 1
    if pages:
        section("Pages", ["Page", "Visuals", "Distinct fields used"])
        for p_, e in pages.items():
            put([p_, len(e["visuals"]), len(e["fields"])])
        row += 1
    if systems:
        section("Data sources (from Power Query)", ["Source system", "Queries", "Query names"])
        for k, v in sorted(systems.items(), key=lambda kv: -len(kv[1])):
            put([k, len(v), ", ".join(sorted(v, key=str.lower))[:300]])
        row += 1
    notes = list(warnings) + ["Auto-generated items (LocalDateTable_*, DateTableTemplate_*, RowNumber-* columns) are excluded."]
    section("Notes", ["Note"])
    for n in notes:
        put([n])
    for i, w in enumerate((34, 70, 30, 12), 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_workbook(data, selected=None, report="", warnings=None, summary=True):
    selected = [n for n in SHEETS if (selected is None or n in selected)]
    wb = Workbook()
    wb.remove(wb.active)
    fill = PatternFill("solid", start_color="1F3864")
    for sheet in selected:
        headers = SHEETS[sheet]
        ws = wb.create_sheet(sheet)
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
            c.fill = fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for row in data[sheet]:
            r = []
            for v in row:
                v = clean(v)
                if isinstance(v, str):
                    v = ILLEGAL_CHARACTERS_RE.sub("", v)
                    if len(v) > 32000:
                        v = v[:32000] + "\n...[truncated]"
                r.append(v)
            ws.append(r)
        for r in ws.iter_rows(min_row=2):
            for c in r:
                c.font = Font(name="Arial", size=10)
                c.alignment = Alignment(vertical="top", wrap_text=isinstance(c.value, str) and len(c.value) > 40)
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.data_type = "s"
        for i, h in enumerate(headers, 1):
            longest = max([len(str(h))] + [min(len(str(r[i - 1])), 70) for r in data[sheet] if i - 1 < len(r)])
            ws.column_dimensions[get_column_letter(i)].width = min(max(11, longest + 2), 70)
        ws.freeze_panes = "A2"
        if data[sheet]:
            ws.auto_filter.ref = ws.dimensions
    if summary:
        _write_summary(wb, data, selected, report, warnings or [])
    if not wb.sheetnames:
        wb.create_sheet("Empty")
    wb.active = 0
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ----------------------------------------------------------------------------
# Local web UI
# ----------------------------------------------------------------------------
PAGE = r"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>PBIX Inspector</title>
<style>
:root{--bg:#f4f6fb;--fg:#1b2233;--card:#fff;--line:#dde2ee;--acc:#2f5bea;--acc2:#1f3864;--mut:#68718a;--ok:#0d8a4f;--bad:#c23b3b;--chip:#eef2ff}
@media(prefers-color-scheme:dark){:root{--bg:#10131b;--fg:#e8ebf3;--card:#1a1f2b;--line:#2c3446;--acc:#6f95ff;--acc2:#27407a;--mut:#98a1b8;--ok:#4cd08b;--bad:#ff7c7c;--chip:#232b3e}}
*{box-sizing:border-box}body{font:14px/1.45 "Segoe UI",system-ui,Arial,sans-serif;background:var(--bg);color:var(--fg);margin:0}
header{padding:18px 28px;background:linear-gradient(90deg,var(--acc2),var(--acc));color:#fff;display:flex;align-items:center;gap:14px}
header h1{font-size:19px;margin:0}header span{opacity:.85;font-size:13px}
main{padding:20px 28px;max-width:1700px;margin:auto}
#drop{border:2px dashed var(--acc);border-radius:14px;padding:34px;text-align:center;background:var(--card);cursor:pointer;transition:.15s}
#drop.over{background:var(--chip);transform:scale(1.005)}#drop b{color:var(--acc)}#drop small{display:block;color:var(--mut);margin-top:6px}
#prog{display:none;margin:16px 0;background:var(--card);border:1px solid var(--line);border-radius:12px;padding:14px 16px}
.bar{height:8px;background:var(--line);border-radius:6px;overflow:hidden;margin-top:8px}.bar i{display:block;height:100%;width:0;background:var(--acc);transition:width .2s}
.bar.ind i{width:35%;animation:slide 1.1s infinite ease-in-out}@keyframes slide{0%{margin-left:-35%}100%{margin-left:100%}}
.msg{margin:12px 0;padding:10px 14px;border-radius:10px;border:1px solid;font-size:13px;white-space:pre-wrap}
.warn{background:#fff7e0;color:#5a4400;border-color:#efd58a}.err{background:#fde8e8;color:#7a1414;border-color:#f0aaaa}
@media(prefers-color-scheme:dark){.warn{background:#3a3217;color:#f2dc9b;border-color:#6b5b22}.err{background:#3b1a1a;color:#ffb4b4;border-color:#7a3030}}
#stats{display:flex;gap:12px;flex-wrap:wrap;margin:16px 0}
.stat{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:10px 16px;min-width:120px;cursor:pointer}
.stat b{display:block;font-size:22px;color:var(--acc)}.stat span{color:var(--mut);font-size:12px}
.bar2{display:flex;gap:10px;align-items:center;flex-wrap:wrap;margin:6px 0 10px}
.tabs{display:flex;gap:6px;flex-wrap:wrap;flex:1}.tab{padding:7px 13px;border:1px solid var(--line);border-radius:9px;background:var(--card);cursor:pointer;font-weight:500}
.tab.on{background:var(--acc);color:#fff;border-color:var(--acc)}.tab em{font-style:normal;opacity:.7;margin-left:5px;font-size:12px}
input#q{padding:8px 12px;border:1px solid var(--line);border-radius:9px;background:var(--card);color:var(--fg);width:300px}
button{border:0;border-radius:9px;padding:9px 16px;cursor:pointer;font-size:14px;font-weight:600}
.dl{background:var(--ok);color:#fff}.ghost{background:var(--chip);color:var(--acc)}
.wrap{overflow:auto;max-height:66vh;border:1px solid var(--line);border-radius:12px;background:var(--card)}
table{border-collapse:separate;border-spacing:0;width:100%}th,td{border-bottom:1px solid var(--line);padding:6px 10px;text-align:left;vertical-align:top;font-size:12.5px}
th{position:sticky;top:0;background:var(--acc2);color:#fff;z-index:2;white-space:nowrap}tr:hover td{background:var(--chip)}
td .c{max-width:460px;white-space:pre-wrap;word-break:break-word;max-height:6.2em;overflow:hidden;cursor:zoom-in}td .c.open{max-height:none;cursor:zoom-out}
.t{color:var(--ok);font-weight:700}.f{color:var(--bad);font-weight:700}
#info{color:var(--mut);font-size:12px;margin:8px 2px}
.selbox{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:12px 16px;margin:0 0 12px}
.selhead,.selfoot{display:flex;gap:10px;align-items:center;flex-wrap:wrap}.sp{flex:1}.mut{color:var(--mut);font-size:12px;flex:1}
#sel{display:flex;flex-wrap:wrap;gap:8px;margin:10px 0}.opt{display:flex;gap:6px;align-items:center;padding:6px 10px;border:1px solid var(--line);border-radius:9px;cursor:pointer;user-select:none}
.opt.on{background:var(--chip);border-color:var(--acc)}.opt small{color:var(--mut)}button:disabled{opacity:.4;cursor:not-allowed}
</style></head><body>
<header><div><h1>PBIX Inspector</h1><span>Model, measures, relationships, M queries, visuals &amp; lineage &middot; runs 100% locally</span></div></header>
<main>
<div id="drop"><div>Drag &amp; drop your <b>.pbix</b> (or .pbit) here, or click to browse</div><small>Nothing is uploaded anywhere &mdash; the file is processed on this computer.</small><input id="f" type="file" accept=".pbix,.pbit" hidden></div>
<div id="prog"><div id="ptxt">Uploading...</div><div class="bar" id="pbar"><i></i></div></div>
<div id="msgs"></div>
<div id="out" hidden>
<div id="stats"></div>
<div class="bar2"><div class="tabs" id="tabs"></div><input id="q" placeholder="Search (space = AND) ..."></div>
<div class="selbox"><div class="selhead"><b>Choose what goes into the Excel file</b> <span id="selcount"></span><span class="sp"></span><button class="ghost" id="sall">Select all</button><button class="ghost" id="snone">Clear</button></div><div id="sel"></div><div class="selfoot"><span class="mut">A Summary sheet is generated automatically as the first tab.</span><button class="dl" id="dl">&#8681; Download selected</button></div></div>
<div class="wrap"><table id="t"></table></div><div id="info"></div>
</div></main>
<script>
const $=id=>document.getElementById(id);let DATA=null,CUR=null,timer=null;
const drop=$('drop'),f=$('f');
drop.onclick=()=>f.click();f.onchange=()=>f.files[0]&&go(f.files[0]);
['dragover','dragenter'].forEach(e=>drop.addEventListener(e,ev=>{ev.preventDefault();drop.classList.add('over')}));
['dragleave','drop'].forEach(e=>drop.addEventListener(e,ev=>{ev.preventDefault();drop.classList.remove('over')}));
drop.addEventListener('drop',ev=>{const x=ev.dataTransfer.files[0];x&&go(x)});
function msg(cls,text){const d=document.createElement('div');d.className='msg '+cls;d.textContent=text;$('msgs').appendChild(d)}
function go(file){
  $('msgs').innerHTML='';$('out').hidden=true;$('prog').style.display='block';
  const bar=$('pbar');bar.className='bar';bar.firstChild.style.width='0';$('ptxt').textContent='Uploading '+file.name+' ...';
  const xhr=new XMLHttpRequest();xhr.open('POST','/analyze');xhr.setRequestHeader('X-Filename',encodeURIComponent(file.name));
  xhr.upload.onprogress=e=>{if(e.lengthComputable)bar.firstChild.style.width=(e.loaded/e.total*100)+'%'};
  xhr.upload.onload=()=>{bar.className='bar ind';const t0=Date.now();clearInterval(timer);
    timer=setInterval(()=>{$('ptxt').textContent='Analyzing model, visuals and lineage ... '+Math.round((Date.now()-t0)/1000)+'s'},300)};
  xhr.onerror=()=>{clearInterval(timer);$('prog').style.display='none';msg('err','Connection to the local tool failed.')};
  xhr.onload=()=>{clearInterval(timer);$('prog').style.display='none';
    let j;try{j=JSON.parse(xhr.responseText)}catch(e){msg('err','Unexpected response: '+xhr.responseText.slice(0,300));return}
    if(j.error){msg('err',j.error);return}
    DATA=j;(j.warnings||[]).forEach(w=>msg('warn',w));
    buildSel(j);
    const st=$('stats');st.innerHTML='';
    j.stats.forEach(s=>{const d=document.createElement('div');d.className='stat';d.innerHTML='<b></b><span></span>';
      d.firstChild.textContent=s[1].toLocaleString();d.lastChild.textContent=s[0];d.onclick=()=>show(s[2]);st.appendChild(d)});
    const tabs=$('tabs');tabs.innerHTML='';
    Object.keys(j.sheets).forEach(n=>{const b=document.createElement('div');b.className='tab';b.dataset.n=n;
      b.innerHTML='<span></span><em></em>';b.firstChild.textContent=n;b.lastChild.textContent=j.sheets[n].total;b.onclick=()=>show(n);tabs.appendChild(b)});
    $('out').hidden=false;show(Object.keys(j.sheets)[0]);};
  xhr.send(file);
}
const SUMMARY='Summary (auto-generated)';let CHECK={};
function buildSel(j){CHECK={};const box=$('sel');box.innerHTML='';
  const names=[SUMMARY].concat(Object.keys(j.sheets));
  names.forEach(n=>{CHECK[n]=true;const l=document.createElement('label');l.className='opt on';
    const c=document.createElement('input');c.type='checkbox';c.checked=true;
    c.onchange=()=>{CHECK[n]=c.checked;l.classList.toggle('on',c.checked);upd()};
    const t=document.createElement('span');t.textContent=n;l.appendChild(c);l.appendChild(t);
    if(n!==SUMMARY){const sm=document.createElement('small');sm.textContent=j.sheets[n].total.toLocaleString()+' rows';l.appendChild(sm)}
    box.appendChild(l)});
  $('sall').onclick=()=>setAll(true);$('snone').onclick=()=>setAll(false);
  $('dl').onclick=()=>{const sh=Object.keys(j.sheets).filter(n=>CHECK[n]);
    location.href='/download/'+j.id+'?sheets='+encodeURIComponent(sh.join(','))+'&summary='+(CHECK[SUMMARY]?1:0)};
  upd();}
function setAll(v){document.querySelectorAll('#sel .opt').forEach(l=>{const c=l.firstChild;c.checked=v;c.onchange()})}
function upd(){const tot=Object.keys(CHECK).length,n=Object.values(CHECK).filter(Boolean).length;
  $('selcount').textContent='('+n+' of '+tot+' selected)';$('dl').disabled=n===0}
function show(n){CUR=n;document.querySelectorAll('.tab').forEach(t=>t.classList.toggle('on',t.dataset.n===n));render()}
$('q').oninput=render;
function render(){if(!DATA)return;const s=DATA.sheets[CUR],terms=$('q').value.toLowerCase().split(/\s+/).filter(Boolean);
  const t=$('t');t.innerHTML='';const h=t.createTHead().insertRow();
  s.headers.forEach(x=>{const th=document.createElement('th');th.textContent=x;h.appendChild(th)});
  const tb=t.createTBody();let shown=0,match=0;
  for(const r of s.rows){
    if(terms.length){const line=r.join(' ').toLowerCase();if(!terms.every(x=>line.includes(x)))continue}
    match++;if(shown>=2000)continue;shown++;
    const tr=tb.insertRow();
    r.forEach(v=>{const td=tr.insertCell();
      if(v===true||v===false){const sp=document.createElement('span');sp.className=v?'t':'f';sp.textContent=v?'TRUE':'FALSE';td.appendChild(sp)}
      else{const d=document.createElement('div');d.className='c';d.textContent=v;d.onclick=()=>d.classList.toggle('open');td.appendChild(d)}})}
  $('info').textContent='Showing '+shown.toLocaleString()+' of '+match.toLocaleString()+' matching rows ('+s.total.toLocaleString()+' total)'+(s.rows.length<s.total?' - preview limited, the Excel file has all rows':'')+(match>shown?' - refine the search to see more':'')+'. Click a long cell to expand it.';
}
</script></body></html>"""

FILES = {}
PREVIEW_ROWS = 8000


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def _send(self, code, body, ctype, extra=None):
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        for k, v in (extra or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        if self.path == "/":
            return self._send(200, PAGE.encode("utf-8"), "text/html; charset=utf-8")
        if self.path.startswith("/download/"):
            u = urlparse(self.path)
            item = FILES.get(u.path.split("/")[-1])
            if not item:
                return self._send(404, b"not found", "text/plain")
            q = parse_qs(u.query)
            chosen = [x for x in (q.get("sheets", [""])[0]).split(",") if x in SHEETS] if "sheets" in q else None
            want_summary = q.get("summary", ["1"])[0] != "0"
            body = build_workbook(item["data"], chosen, item["report"], item["warnings"], want_summary)
            return self._send(200, body, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                              {"Content-Disposition": f'attachment; filename="{item["report"]}_inspection.xlsx"'})
        self._send(404, b"not found", "text/plain")

    def do_POST(self):
        if self.path != "/analyze":
            return self._send(404, b"not found", "text/plain")
        tmp = None
        try:
            fname = unquote(self.headers.get("X-Filename", "model.pbix"))
            length = int(self.headers.get("Content-Length", 0))
            fd, tmp = tempfile.mkstemp(suffix=os.path.splitext(fname)[1] or ".pbix")
            with os.fdopen(fd, "wb") as f:
                left = length
                while left > 0:
                    chunk = self.rfile.read(min(1 << 20, left))
                    if not chunk:
                        break
                    f.write(chunk)
                    left -= len(chunk)
            base = os.path.splitext(os.path.basename(fname))[0]
            data, warnings = inspect(tmp, base)
            fid = uuid.uuid4().hex
            FILES[fid] = {"data": data, "report": base, "warnings": warnings}
            sheets = {n: {"headers": SHEETS[n], "total": len(data[n]),
                          "rows": [[clean(v) for v in r] for r in data[n][:PREVIEW_ROWS]]} for n in SHEETS}
            nvis = len({(v["page"], v["vid"]) for v in data["_visuals"] if v["vtype"] not in ("Page Filters", "Report Filters")})
            stats = [["Tables", len(data["Tables"]), "Tables"], ["Columns", len(data["Columns"]), "Columns"],
                     ["Measures", len(data["Measures"]), "Measures"],
                     ["Relationships", len(data["Relationships"]), "Relationships"],
                     ["M Queries", len(data["M Queries"]), "M Queries"],
                     ["Visuals", nvis, "Visual Design"], ["Lineage rows", len(data["Lineage"]), "Lineage"],
                     ["Impact rows", len(data["Impact Analysis"]), "Impact Analysis"]]
            self._send(200, json.dumps({"id": fid, "sheets": sheets, "warnings": warnings, "stats": stats},
                                       default=str).encode("utf-8"), "application/json")
        except Exception as e:
            self._send(200, json.dumps({"error": f"{type(e).__name__}: {e}"}).encode("utf-8"), "application/json")
        finally:
            if tmp and os.path.exists(tmp):
                os.remove(tmp)


def serve():
    port = 8765
    while True:
        try:
            srv = ThreadingHTTPServer(("127.0.0.1", port), Handler)
            break
        except OSError:
            port += 1
    url = f"http://127.0.0.1:{port}/"
    print(f"PBIX Inspector running at {url}  (Ctrl+C to stop)")
    threading.Timer(0.6, lambda: webbrowser.open(url)).start()
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        pass


def _pause():
    """Keep the console open when launched by double-click / drag onto the .exe."""
    if getattr(sys, "frozen", False):
        try:
            input("\nPress Enter to close this window...")
        except Exception:
            pass


def main():
    try:
        import multiprocessing
        multiprocessing.freeze_support()
    except Exception:
        pass
    try:
        if len(sys.argv) > 1:
            path = sys.argv[1].strip('"')
            data, warnings = inspect(path)
            chosen = [x.strip() for x in sys.argv[2].split(",")] if len(sys.argv) > 2 else None
            out_path = os.path.splitext(path)[0] + "_inspection.xlsx"
            with open(out_path, "wb") as f:
                f.write(build_workbook(data, chosen, os.path.splitext(os.path.basename(path))[0], warnings))
            for w in warnings:
                print("WARNING:", w)
            print("Saved", out_path)
            _pause()
        else:
            print("Keep this window open while you use the tool. Close it to stop.")
            serve()
    except Exception:
        import traceback
        traceback.print_exc()
        _pause()


if __name__ == "__main__":
    main()